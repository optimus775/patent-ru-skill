#!/usr/bin/env python3
"""
Конвертирует PowerPoint (.pptx/.ppsx) в Markdown и извлекает встроенные изображения для Step 2.

Зависимость: python-pptx; см. requirements.txt в корне репозитория.

Примеры:
  python pptx_to_md.py --input review.pptx --output outputs/case/review.md
  python pptx_to_md.py -i a.pptx -o b/out.md --media-dir b/slide_images

Каталог изображений по умолчанию: соседний каталог {md_stem}_media/.
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path


def _require_pptx():
    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
    except ImportError:
        print(
            "Не найдена зависимость python-pptx. Выполните в корне skill: pip install -r requirements.txt",
            file=sys.stderr,
        )
        sys.exit(1)
    return Presentation, MSO_SHAPE_TYPE


def _walk_shapes(shapes, MSO_SHAPE_TYPE):
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _walk_shapes(shape.shapes, MSO_SHAPE_TYPE)
        else:
            yield shape


def _shape_text(shape) -> str:
    if getattr(shape, "has_text_frame", False):
        t = (shape.text_frame.text or "").strip()
        return t
    if getattr(shape, "has_table", False):
        rows = []
        for row in shape.table.rows:
            cells = []
            for cell in row.cells:
                cells.append((cell.text or "").strip().replace("\n", " "))
            rows.append("| " + " | ".join(cells) + " |")
        if rows:
            return "\n".join(rows)
    return ""


def _rel_media_path(out_file: Path, media_file: Path) -> str:
    try:
        return media_file.relative_to(out_file.parent).as_posix()
    except ValueError:
        return media_file.as_posix()


def _run(input_pptx: Path, output_md: Path, media_dir: Path | None) -> int:
    Presentation, MSO_SHAPE_TYPE = _require_pptx()

    if not input_pptx.is_file():
        print(f"Входной файл не найден: {input_pptx}", file=sys.stderr)
        return 2
    suf = input_pptx.suffix.lower()
    if suf not in (".pptx", ".ppsx"):
        print("Предупреждение: ожидается .pptx / .ppsx (OOXML); старый .ppt не поддерживается.", file=sys.stderr)

    output_md = output_md.resolve()
    output_md.parent.mkdir(parents=True, exist_ok=True)

    if media_dir is None:
        media_dir = output_md.parent / f"{output_md.stem}_media"
    else:
        media_dir = media_dir.resolve()
    media_dir.mkdir(parents=True, exist_ok=True)

    try:
        prs = Presentation(str(input_pptx))
    except Exception as e:
        print(f"Не удалось открыть презентацию: {e}", file=sys.stderr)
        return 3

    lines: list[str] = [
        f"<!-- Конвертировано pptx_to_md.py из {input_pptx.name}; не редактируйте эту строку метаданных вручную. -->\n"
    ]
    img_counter = [0]

    for sn, slide in enumerate(prs.slides, start=1):
        lines.append(f"\n## Слайд {sn}\n")

        for shape in _walk_shapes(slide.shapes, MSO_SHAPE_TYPE):
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    img = shape.image
                    ext = (img.ext or "png").lower()
                    if ext == "jpeg":
                        ext = "jpg"
                    img_counter[0] += 1
                    fname = f"slide{sn:02d}_img{img_counter[0]:04d}.{ext}"
                    out_img = media_dir / fname
                    out_img.write_bytes(img.blob)
                    rel = _rel_media_path(output_md, out_img)
                    lines.append(f"\n![]({rel})\n")
                except Exception as e:
                    print(f"Предупреждение: не удалось извлечь изображение со слайда {sn}: {e}", file=sys.stderr)
                continue

            block = _shape_text(shape)
            if block:
                lines.append(block)
                lines.append("\n\n")

        try:
            nf = slide.notes_slide.notes_text_frame
            note_txt = (nf.text or "").strip() if nf is not None else ""
            if note_txt:
                lines.append("\n**Заметки**:\n\n")
                lines.append(note_txt)
                lines.append("\n\n")
        except (AttributeError, ValueError):
            pass

    body = "".join(lines).rstrip() + "\n"
    output_md.write_text(body, encoding="utf-8")

    print(f"Записано: {output_md}")
    print(f"Каталог изображений: {media_dir}")
    return 0


def main() -> int:
    p = argparse.ArgumentParser(description="PowerPoint (.pptx/.ppsx) -> Markdown с извлечением изображений")
    p.add_argument("-i", "--input", required=True, type=Path, help="Путь к входному .pptx/.ppsx")
    p.add_argument("-o", "--output", required=True, type=Path, help="Путь к выходному .md")
    p.add_argument(
        "--media-dir",
        type=Path,
        default=None,
        help="Каталог для изображений; по умолчанию {md_stem}_media рядом с .md",
    )
    args = p.parse_args()
    return _run(args.input, args.output, args.media_dir)


if __name__ == "__main__":
    raise SystemExit(main())
